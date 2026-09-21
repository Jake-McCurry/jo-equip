#!/usr/bin/env python3
"""Offline importer/checker for the approved category metadata workbook."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT.parents[1] / "attached_assets" / "SEO_Titles___Meta_Descr._1790014637209.xlsx"
DEFAULT_OUTPUT = ROOT / "src" / "data" / "categorySeo.json"
EXPECTED_ROWS = 1264
EXPECTED_ORIGIN = "https://equip.jesusonline.com"


def import_workbook(path: Path) -> dict[str, dict[str, str]]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    metadata: dict[str, dict[str, str]] = {}
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        url, title, description = row[1:4]
        if not all(isinstance(value, str) and value for value in (url, title, description)):
            raise ValueError(f"Row {row_number} is missing URL, title, or description")
        parsed = urlsplit(url)
        origin = f"{parsed.scheme}://{parsed.netloc}"
        if origin != EXPECTED_ORIGIN:
            raise ValueError(f"Row {row_number} has unexpected origin: {url}")
        path_key = parsed.path.rstrip("/") or "/"
        if path_key in metadata:
            raise ValueError(f"Row {row_number} duplicates URL path: {path_key}")
        metadata[path_key] = {"title": title, "description": description}
    if len(metadata) != EXPECTED_ROWS:
        raise ValueError(f"Expected {EXPECTED_ROWS} rows, found {len(metadata)}")
    return metadata


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()

    imported = import_workbook(args.workbook)
    serialized = json.dumps(imported, ensure_ascii=False, indent=2) + "\n"
    if args.check:
        checked_in = args.output.read_text(encoding="utf-8")
        if checked_in != serialized:
            raise SystemExit(f"{args.output} does not exactly match {args.workbook}")
        print(f"Validated {len(imported)} exact metadata paths against workbook")
        return
    args.output.write_text(serialized, encoding="utf-8")
    print(f"Wrote {len(imported)} exact metadata paths to {args.output}")


if __name__ == "__main__":
    main()