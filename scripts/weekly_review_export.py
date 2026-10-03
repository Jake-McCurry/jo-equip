"""Export a read-only checkpoint review; does not change application files."""
import json, re, subprocess
from datetime import datetime
from zoneinfo import ZoneInfo
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

def git(*args):
    return subprocess.check_output(["git", *args], text=True)

start = "2026-09-27T00:00:00-07:00"
end = "2026-10-04T00:00:00-07:00"
commits = [x.split("|", 2) for x in git("log", "--reverse", "--since="+start, "--until="+end,
    "--format=%H|%aI|%s").splitlines()]
wb = Workbook()
intro = wb.active
intro.title = "Read Me"
intro.append(["Weekly change review", "September 27–October 3, 2026 (America/Los_Angeles)"])
notes = [
("Scope", "All available local checkpoint history within the seven-day window, plus current conversation verification results. Checkpoint timestamps are evidence dates, not necessarily the precise implementation time."),
("How to review", "Filter the Review sheet by area. Select Pending, Approved, Needs changes, or Not applicable, and enter your notes. Metadata and redirects have separate rows for individual review."),
("Coverage", "The File History sheet inventories every changed file per checkpoint, including uploads, generated files and maintenance. These are not all independent user-facing changes."),
("Publication", "Implemented means present in the workspace history. No production deployment was verified for this report."),
("Limits", "Local history and retained conversation are the available evidence; this is not a guarantee that every action in other sessions or external services is recorded. Dates are not inferred from uploaded-file timestamps."),
("Other-project files", "The FAQ replacement/removal request was withdrawn as belonging to another project. The Guide formatting request is on hold. Neither is counted as implemented."),
("No-op checks", "SEO-018 and SEO-020 were verified as already implemented. Verification is listed separately to avoid counting it as a new change."),
("Exact details", "Code Details contains per-file patches for non-generated application source and scripts, including old/new wording. Large patches are split into numbered parts. File History covers binary assets."),
]
for n in notes: intro.append(n)
review = wb.create_sheet("Review")
review.append(["ID", "Checkpoint date (Pacific)", "Area", "Page / item", "Change", "Evidence", "Status", "Publication", "Review status", "Your notes"])
history = wb.create_sheet("File History")
history.append(["Date (Pacific)", "Checkpoint", "Checkpoint description", "File", "Added lines", "Removed lines", "Classification"])
details = wb.create_sheet("Code Details")
details.append(["Date (Pacific)", "Checkpoint", "File", "Part", "Exact patch"])
def add(date, area, item, change, evidence, status="Implemented"):
    review.append([f"R{review.max_row:03}",date,area,item,change,evidence,status,"Not verified","Pending",""])

def category(path):
    if path.startswith("attached_assets/"): return "Reference upload—not implementation"
    if "/generated/" in path or "sitemap" in path or "/.astro/" in path: return "Generated / supporting output"
    if path.endswith(".pdf"): return "PDF asset"
    if "/memory/" in path or path.endswith(".md"): return "Documentation"
    return "Code / other asset"

prefix = "artifacts/discipleship-hub/"
for sha, timestamp, subject in commits:
    date = datetime.fromisoformat(timestamp).astimezone(ZoneInfo("America/Los_Angeles")).strftime("%Y-%m-%d %H:%M")
    stats = git("diff-tree","--root","--no-commit-id","--numstat","-r",sha).splitlines()
    for line in stats:
        parts = line.split("\t",2)
        if len(parts)!=3: continue
        plus,minus,path=parts
        history.append([date,sha[:7],subject,path,plus,minus,category(path)])
        source = (path.startswith(("artifacts/","scripts/")) and "/generated/" not in path
                  and "/.astro/" not in path and path.endswith((".astro",".ts",".tsx",".js",".mjs",".css",".json"))
                  and not any(s in path for s in ["package-lock","bible/data","public/","pnpm-lock"]))
        if not source: continue
        patch = git("show",sha,"--format=","--",path)
        for i in range(0,len(patch),28000):
            details.append([date,sha[:7],path,i//28000+1,patch[i:i+28000]])
        if path.endswith("categorySeoOverrides.json"):
            try: old=json.loads(git("show",sha+"^:"+path))
            except Exception: old={}
            new=json.loads(git("show",sha+":"+path))
            for url, record in new.items():
                for key,value in record.items():
                    if old.get(url,{}).get(key)!=value:
                        add(date,"SEO metadata / headings",url,
                            f"{key}: {value}\nPrevious override: {old.get(url,{}).get(key,'No override (inherited)')}",sha[:7]+" / "+path)
    # Checkpoint-specific user-facing changes not duplicated by the generic file audit.
    short=sha[:7]
    items={
      "612902f":[("Bible reader","/bible","Added the NET Bible reader and header navigation entry"),
                 ("Bible reader","Passages","Book/chapter navigation and verse passage display"),
                 ("Bible reader","Study tools","Verse study panel; local notes/highlights/bookmarks and study library"),
                 ("Bible reader","Backup and recap","Study-data backup/import and recap interface"),
                 ("Bible reader","Delivery and attribution","Bible catalog/data endpoints and translation attribution")],
      "50fdc10":[("Leader kits","/leader-kits","Updated available kit links and labels")],
      "2ccaf77":[("Leader kits","Kit covers","Updated covers and added repeatable PDF-cover build script")],
      "eb8d117":[("Leader kits","Signup campaign variants","Sunday campaign tagging and button wording; keep OpenAI campaign untagged"),
                 ("Leader kits","Signup feedback","Restore the correct campaign-specific button label after submission")],
      "d2b6950":[("External links","Shared layout","Added external hyperlink nofollow policy while preserving internal links")],
      "2704f1b":[("External links","Dynamic content","Extended external-link handling"),("Brand assets","Site icons","Updated favicon / touch-icon assets and layout references")],
      "bb1a2e0":[("Category pages","Category template","Removed temporary category-level leader-kit promotional block")],
      "54402c7":[("Bible Study Tools","Article contents","Added responsive article contents navigation"),
                 ("Bible Study Tools","Article additions","Added ministry-needs, benefits, and FAQ sections; web-only scope")],
      "f5a49a0":[("Bible Study Tools","Contents and FAQs","Refined contents labels, FAQ wording and accordion formatting")],
      "2957dac":[("Duplicate articles","Routing and links","Added exact retirement manifest, 301 handling, survivor links and sitemap filtering; individual pairs below")],
      "37f74c4":[("Duplicate articles","Majesty editorial decision","Retained six approved local Attributes versions; retired corresponding Majesty duplicates")],
    }
    for area,item,change in items.get(short,[]): add(date,area,item,change,short)
    if short=="d2c63eb":
        for url in ["pastors","home-church-leaders","chaplains","missionaries"]:
            add(date,"SEO-020 schema","/"+url,"Added reusable WebPage and Home → current page BreadcrumbList JSON-LD",short)
    if short=="9cc7119":
        text=Path("attached_assets/Pasted--SEO-Task-SEO-018-Task-Name-Update-Meta-Titles-Meta-Des_1790808920553.txt").read_text()
        for url,title,desc in re.findall(r'^\|\s*(https://equip\.jesusonline\.com\S*)\s*\|\s*\*\*(.*?)\*\*\s*\|\s*\*\*(.*?)\*\*\s*\|',text,re.M):
            if url.endswith("/categories/church"): continue
            add(date,"SEO-018 metadata",url,f"Title: {title}\nDescription: {desc}",short)
    if short=="2386811":
        for path in ["src/pages/books.astro","src/pages/books/[id].astro"]:
            patch=git("show",sha,"--format=","--",prefix+path)
            add(date,"Book metadata / heading",path,"Updated page-specific SEO copy"+("; Books heading changed to Christian Books Online" if "[id]" not in path else " for A Heart After God only"),short+"; exact wording in Code Details")
    if short=="a3214a1":
        text=Path("attached_assets/Pasted-S-No-From-To-Anchor-Sentence-Section-1-https-equip-jesu_1790806641492.txt").read_text()
        # Preserve individual specification rows verbatim rather than paraphrasing anchor text.
        for line in text.splitlines():
            if re.match(r'^\s*\d+\s+https',line):
                number,source,target,anchor,sentence,section=line.split("\t",5)
                add(date,"Internal linking",source,
                    f'Link {number}: "{anchor}" → {target}\nSection: {section}\nSentence: {sentence}',
                    short+"; attachment + Code Details")
        if not any(r[2].value=="Internal linking" for r in list(review.rows)[1:]):
            add(date,"Internal linking","11 approved opportunities","Implemented 11 opportunities, including seven repeated FAQ passages. Exact additions in Code Details.",short)
    # Binary PDFs get their own review rows.
    for line in stats:
        p=line.split("\t",2)[-1]
        if p.startswith("artifacts/") and p.endswith(".pdf"):
            add(date,"PDF resources",p,"Added or updated PDF asset (binary; inspect file visually)",short)

redirects=json.loads(subprocess.check_output(["node","--input-type=module","-e",
 "import {ARTICLE_REDIRECTS} from './artifacts/discipleship-hub/src/data/articleCanonicalPaths.mjs';console.log(JSON.stringify(ARTICLE_REDIRECTS))"],text=True))
for source,target in redirects.items():
    add("2026-09-30","Individual article redirect",source,"Retire duplicate; 301 redirect to "+target,"articleCanonicalPaths.mjs; 2957dac / 37f74c4")
for task,scope in [("SEO-018","Nine approved metadata pages"),("SEO-020","Four persona pages: WebPage + BreadcrumbList")]:
    add("Current conversation","Verification only",task,scope+" already implemented; rendered development HTML checked, no new edits.","Conversation verification","Verified—no change")
add("Current conversation","On hold","The Guide format refinements","Spreadsheet and two table images received; no implementation.","User: hold on that","On hold")
add("Current conversation","Withdrawn","FAQ and Go Deeper removals/replacement","Belongs to a different project; no implementation.","User withdrew request","Withdrawn")

for idx,ws in enumerate(wb):
    ws.freeze_panes="A2"
    ws.sheet_view.showGridLines=False
    ws.row_dimensions[1].height=30
    for c in ws[1]:
        c.fill=PatternFill("solid",fgColor="194C69")
        c.font=Font(color="FFFFFF",bold=True)
        c.alignment=Alignment(wrap_text=True,vertical="center")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment=Alignment(vertical="top",wrap_text=True)
        ws.row_dimensions[row[0].row].height=80 if ws.title!="Code Details" else 120
    widths={"Read Me":[27,120],"Review":[10,22,28,58,100,50,24,22,22,55],
            "File History":[22,15,55,100,15,15,35],"Code Details":[22,15,85,9,145]}[ws.title]
    for i,width in enumerate(widths,1):
        ws.column_dimensions[ws.cell(1,i).column_letter].width=width
    if ws.title!="Read Me":
        t=Table(displayName=f"ReviewTable{idx}",ref=ws.dimensions)
        t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True)
        ws.add_table(t)
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation="landscape"
    ws.page_setup.fitToWidth=1
    ws.page_setup.fitToHeight=0
    ws.print_title_rows="1:1"
dv=DataValidation(type="list",formula1='"Pending,Approved,Needs changes,Not applicable"')
review.add_data_validation(dv)
dv.add(f"I2:I{review.max_row}")
for row in review.iter_rows(min_row=2,min_col=9,max_col=10):
    for c in row: c.fill=PatternFill("solid",fgColor="FFF2CC")
out=Path("outputs/weekly-review/JO-EQUIP-Weekly-Changes-2026-09-27-to-2026-10-03.xlsx")
out.parent.mkdir(parents=True,exist_ok=True)
wb.save(out)
check=load_workbook(out)
assert check["Review"].max_row==review.max_row
assert len(redirects)==65
print(f"{out}: {review.max_row-1} review rows; {history.max_row-1} file history rows; {details.max_row-1} patch entries; 65 individual redirects.")